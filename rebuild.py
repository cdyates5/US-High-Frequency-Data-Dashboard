#!/usr/bin/env python3
"""
Acheron Insights — US High-Frequency Economic Data dashboard
Data build pipeline. Fetches all sources, computes seasonal comparisons,
and regenerates the single-file HTML dashboard in place.

Sources (all public):
  FRED (api.stlouisfed.org)           WEI, GDP proxy series, ICSA, CCSA, TOTLL,
                                       NFCI, MORTGAGE30US, STLFSI4, rail carloads/intermodal
  Atlanta Fed GDPNow xlsx             daily topline nowcast
  US Census BFS weekly (NSA)          weekly business applications
  TSA passenger volumes               daily throughput (official + validated archive)
  Cleveland Fed inflation nowcast     monthly CPI nowcast (daily-updated)
"""
import requests, json, io, os, sys, datetime as dt
from collections import OrderedDict, defaultdict

HDR = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"}
FRED_KEY = os.environ.get("FRED_API_KEY")
if not FRED_KEY:
    sys.exit(
        "Missing FRED_API_KEY.\n"
        "Get a free key (instant, no approval wait) at:\n"
        "  https://fred.stlouisfed.org/docs/api/api_key.html\n"
        "Then run:\n"
        "  export FRED_API_KEY=your_key_here      # macOS/Linux\n"
        "  setx FRED_API_KEY your_key_here         # Windows\n"
    )
TIMEOUT = 60

def fred(series, start=None):
    url = (f"https://api.stlouisfed.org/fred/series/observations"
           f"?series_id={series}&api_key={FRED_KEY}&file_type=json")
    if start:
        url += f"&observation_start={start}"
    r = requests.get(url, headers=HDR, timeout=TIMEOUT)
    r.raise_for_status()
    out = []
    for o in r.json().get("observations", []):
        if o["value"] not in (".", ""):
            out.append([o["date"], float(o["value"])])
    return out

def yoy(series_pairs, periods):
    """Year-over-year % change given ~weekly data (periods = 52)."""
    vals = series_pairs
    out = []
    for i in range(periods, len(vals)):
        prev = vals[i - periods][1]
        if prev:
            out.append([vals[i][0], (vals[i][1] / prev - 1.0) * 100.0])
    return out

def sma(pairs, w):
    """Centered simple moving average over a list of [date, value]."""
    v = [p[1] for p in pairs]
    n = len(v); half = w // 2
    out = []
    for i in range(n):
        lo = max(0, i - half); hi = min(n, i + half + 1)
        seg = v[lo:hi]
        out.append([pairs[i][0], sum(seg) / len(seg)])
    return out

DATA = {}

# ---------------------------------------------------------------- FRED weekly / monthly
print("FRED WEI...");            DATA["wei"]      = fred("WEI", "2021-01-01")
print("FRED ICSA...");          DATA["icsa"]     = [[d, v/1000] for d, v in fred("ICSA", "2022-01-01")]
print("FRED CCSA...");          DATA["ccsa"]     = [[d, v/1000] for d, v in fred("CCSA", "2022-01-01")]
print("FRED NFCI...");          DATA["nfci"]     = fred("NFCI", "2021-01-01")
print("FRED MORTGAGE30US...");  DATA["mortgage"] = fred("MORTGAGE30US", "2020-01-01")
print("FRED STLFSI4...");       DATA["stress"]   = fred("STLFSI4", "2021-01-01")

print("FRED TOTLL (loan growth)...")
totll = fred("TOTLL")                       # weekly, need full history for YoY
DATA["loan_yoy"] = [p for p in yoy(totll, 52) if p[0] >= "2018-01-01"]

print("FRED rail carloads + intermodal...")
carloads   = dict(fred("RAILFRTCARLOADSD11", "2020-01-01"))
intermodal = dict(fred("RAILFRTINTERMODALD11", "2020-01-01"))
rail = []
for d in sorted(set(carloads) & set(intermodal)):
    rail.append([d, (carloads[d] + intermodal[d]) / 1000.0])   # thousands of units
DATA["rail"] = rail

# Seasonal view needs NON-seasonally-adjusted data (the SA series above has the
# seasonal signal removed by construction). Pull NSA carloads + intermodal,
# combine, and structure by calendar month: 2021-2025 mean / min / max vs 2026.
print("FRED rail NSA (seasonal)...")
car_nsa   = dict(fred("RAILFRTCARLOADS",   "2020-01-01"))
inter_nsa = dict(fred("RAILFRTINTERMODAL", "2020-01-01"))
rail_nsa = {d: car_nsa[d] + inter_nsa[d]                # actual units (carloads + intermodal)
            for d in set(car_nsa) & set(inter_nsa)}
rail_bymonth = defaultdict(dict)                       # month(1-12) -> {year: value}
for d, v in rail_nsa.items():
    rail_bymonth[int(d[5:7])][int(d[:4])] = v
r_mean, r_lo, r_hi, r_cur = [], [], [], []
cur_year = 2026
_base = [2021, 2022, 2023, 2024, 2025]
for m in range(1, 13):
    yr_vals = [rail_bymonth[m][y] for y in _base if y in rail_bymonth[m]]
    if yr_vals:
        r_mean.append(round(sum(yr_vals)/len(yr_vals))); r_lo.append(round(min(yr_vals))); r_hi.append(round(max(yr_vals)))
    else:
        r_mean.append(None); r_lo.append(None); r_hi.append(None)
    cv = rail_bymonth[m].get(cur_year)
    r_cur.append(round(cv) if cv is not None else None)
_cur_months = [m for m in range(1, 13) if rail_bymonth[m].get(cur_year) is not None]
DATA["rail_seasonal"] = {
    "months": list(range(1, 13)),
    "mean": r_mean, "lo": r_lo, "hi": r_hi, "current": r_cur,
    "base_label": "2021–2025 avg",
    "last_month": max(_cur_months) if _cur_months else None,
}

# ---------------------------------------------------------------- Atlanta Fed GDPNow (daily)
print("Atlanta Fed GDPNow xlsx...")
import openpyxl
u = ("https://www.atlantafed.org/-/media/Project/Atlanta/FRBA/Documents/"
     "cqer/researchcq/gdpnow/GDPTrackingModelDataAndForecasts.xlsx")
xb = requests.get(u, headers=HDR, timeout=120).content
wb = openpyxl.load_workbook(io.BytesIO(xb), read_only=True, data_only=True)
pairs = []
# historical daily archive
for r in wb["TrackingArchives"].iter_rows(values_only=True):
    d, g = r[0], r[27]
    if isinstance(d, dt.datetime) and g is not None:
        try: pairs.append((d.date(), float(g)))
        except: pass
# current-quarter evolution (topline GDP*)
ws = wb["CurrentQtrEvolution"]; rows = list(ws.iter_rows(values_only=True))
hdr = rows[0]; blocks = [j for j in range(len(hdr)) if hdr[j] == "Date"]
for b in blocks:
    for r in rows[1:]:
        d, g = r[b], (r[b+2] if b+2 < len(hdr) else None)
        if isinstance(d, dt.datetime) and g is not None:
            try: pairs.append((d.date(), float(g)))
            except: pass
od = OrderedDict()
for d, g in sorted(pairs): od[d] = g
DATA["gdpnow"] = [[d.isoformat(), round(g, 3)] for d, g in od.items() if d >= dt.date(2022, 1, 1)]

# ---------------------------------------------------------------- Cleveland Fed CPI nowcast
print("Cleveland Fed inflation nowcast...")
j = requests.get("https://www.clevelandfed.org/-/media/files/webcharts/"
                 "inflationnowcasting/nowcast_month.json?sc_lang=en",
                 headers=HDR, timeout=TIMEOUT).json()
mom = {}   # target-month -> headline CPI nowcast MoM %
for entry in j:
    ch = entry.get("chart", {}); sub = ch.get("subcaption", "")
    try:
        y, m = sub.split("-"); key = (int(y), int(m))
    except Exception:
        continue
    cpi_series = None
    for s in entry.get("dataset", []):
        nm = (s.get("seriesname") or "").lower()
        if nm == "cpi inflation":            # headline, not core, not "actual"
            cpi_series = s; break
    if cpi_series is None:
        continue
    vals = [d.get("value") for d in cpi_series.get("data", []) if d.get("value") not in (None, "")]
    if vals:
        try: mom[key] = float(vals[-1])       # final (current) MoM nowcast for that month
        except: pass
# chain 12 consecutive monthly MoM nowcasts into a YoY inflation rate
def add_months(y, m, k):
    idx = (y*12 + (m-1)) + k; return idx//12, idx%12 + 1
cpi_yoy = []
for (y, m) in sorted(mom.keys()):
    window = [mom.get(add_months(y, m, -k)) for k in range(0, 12)]
    if all(v is not None for v in window):
        prod = 1.0
        for v in window: prod *= (1.0 + v/100.0)
        cpi_yoy.append([f"{y:04d}-{m:02d}-01", round((prod - 1.0) * 100.0, 2)])
DATA["cpi_nowcast"] = [p for p in cpi_yoy if p[0] >= "2021-01-01"]

# ---------------------------------------------------------------- SEASONAL: TSA (daily)
print("TSA daily (archive + official)...")
csv = {}
r = requests.get("https://raw.githubusercontent.com/DavidTeju/flight-statistics/"
                 "HEAD/static/tsa_passenger_volumes.csv", headers=HDR, timeout=TIMEOUT)
for line in r.text.strip().splitlines()[1:]:
    d, v = line.split(","); csv[d.strip()] = int(v)
# overlay official live page (validated identical; extends to latest)
import re
live_html = requests.get("https://www.tsa.gov/travel/passenger-volumes", headers=HDR, timeout=TIMEOUT).text
for row in re.findall(r"<tr[^>]*>(.*?)</tr>", live_html, re.S)[1:]:
    cells = [re.sub(r"<[^>]+>", "", c).strip() for c in re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", row, re.S)]
    if len(cells) >= 2 and "/" in cells[0]:
        mm, dd, yy = cells[0].split("/")
        try: csv[f"{yy}-{int(mm):02d}-{int(dd):02d}"] = int(cells[1].replace(",", ""))
        except: pass

def build_seasonal_daily(daily, cur_year, base_years):
    """Return {mmdd:[...], current:[...], mean:[...], lo:[...], hi:[...]} aligned Jan1->Dec31."""
    by_md = defaultdict(dict)      # mmdd -> {year: value}
    cur = {}
    for ds, val in daily.items():
        y = int(ds[:4]); md = ds[5:]
        if md == "02-29":
            continue
        by_md[md][y] = val
        if y == cur_year:
            cur[md] = val
    md_axis = sorted(by_md.keys())
    mean, lo, hi, cur_line = [], [], [], []
    for md in md_axis:
        yr_vals = [by_md[md][y] for y in base_years if y in by_md[md]]
        if yr_vals:
            mean.append(sum(yr_vals)/len(yr_vals)); lo.append(min(yr_vals)); hi.append(max(yr_vals))
        else:
            mean.append(None); lo.append(None); hi.append(None)
        cur_line.append(cur.get(md))
    return md_axis, mean, lo, hi, cur_line

def smooth_list(x, w=7, mask=True):
    """Centered MA. When mask=True, keep None wherever the raw value is None
    (prevents the current-year line from bleeding past its last real date)."""
    out = []; half = w // 2
    for i in range(len(x)):
        if mask and x[i] is None:
            out.append(None); continue
        seg = [v for v in x[max(0,i-half):min(len(x),i+half+1)] if v is not None]
        out.append(round(sum(seg)/len(seg), 1) if seg else None)
    return out

md_axis, mean, lo, hi, cur = build_seasonal_daily(csv, 2026, [2021,2022,2023,2024,2025])
DATA["tsa_seasonal"] = {
    "axis": md_axis,
    "mean": smooth_list(mean), "lo": smooth_list(lo), "hi": smooth_list(hi),
    "current": smooth_list(cur, mask=True),
    "base_label": "2021–2025 avg",
    "last_date": max(csv.keys()),
}

# ---------------------------------------------------------------- SEASONAL: Business apps (weekly)
print("Census BFS weekly business applications...")
r = requests.get("https://www.census.gov/econ/bfs/csv/bfs_us_apps_weekly_nsa.csv", headers=HDR, timeout=TIMEOUT)
rows = [l.split(",") for l in r.text.strip().splitlines()[1:]]
# columns: Year,Week,BA_NSA,...
ba = defaultdict(dict)   # week -> {year: BA_NSA}
cur_ba = {}
for row in rows:
    try:
        y = int(row[0]); wk = int(row[1]); val = float(row[2])
    except (ValueError, IndexError):
        continue
    if wk >= 53: continue
    ba[wk][y] = val
    if y == 2026: cur_ba[wk] = val
# date_table for week -> month label
dt_tab = requests.get("https://www.census.gov/econ/bfs/csv/date_table.csv", headers=HDR, timeout=TIMEOUT).text
wk_month = {}
for l in dt_tab.strip().splitlines()[1:]:
    c = l.split(",")
    try:
        y = int(c[0]); wk = int(c[1])
        if y == 2025:   # reference calendar for labels
            mm = int(c[2].split("/")[0]); wk_month[wk] = mm
    except: pass
weeks = sorted(ba.keys())
base_years = [2021,2022,2023,2024,2025]
b_mean, b_lo, b_hi, b_cur, months = [], [], [], [], []
for wk in weeks:
    yr_vals = [ba[wk][y] for y in base_years if y in ba[wk]]
    if yr_vals:
        b_mean.append(round(sum(yr_vals)/len(yr_vals))); b_lo.append(round(min(yr_vals))); b_hi.append(round(max(yr_vals)))
    else:
        b_mean.append(None); b_lo.append(None); b_hi.append(None)
    b_cur.append(round(cur_ba[wk]) if wk in cur_ba else None)
    months.append(wk_month.get(wk, 0))
DATA["bizapps_seasonal"] = {
    "weeks": weeks, "months": months,
    "mean": b_mean, "lo": b_lo, "hi": b_hi, "current": b_cur,
    "base_label": "2021–2025 avg",
    "last_week": max(cur_ba.keys()) if cur_ba else None,
}

# ------------------------------------------------- initial claims (seasonal)
# Seasonal comparison needs the NON-seasonally-adjusted claims series (ICNSA);
# the indicator-grid panel uses ICSA (adjusted). Weekly, keyed by week-of-year
# so the same slot ~= same calendar time across years. NOTE: the 2021-2025
# baseline includes 2021's pandemic-elevated claims, which widens the upper
# band in Q1 and pulls the average up — kept for consistency with the other
# seasonal panels.
print("FRED ICNSA (initial claims, seasonal)...")
icnsa = fred("ICNSA", "2021-01-01")
def week_of_year(ds):
    d = dt.date.fromisoformat(ds)
    return min(52, (d.timetuple().tm_yday - 1) // 7 + 1)
clm = defaultdict(dict)                                  # week -> {year: value}
cur_clm, cur_clm_date = {}, {}
for ds, v in icnsa:
    y, w = int(ds[:4]), week_of_year(ds)
    if y == 2026:
        if w not in cur_clm or ds > cur_clm_date[w]:      # keep latest obs in a week
            cur_clm[w] = v; cur_clm_date[w] = ds
    elif 2021 <= y <= 2025:
        clm[w][y] = v
# month label per week (from a non-leap reference year)
def week_month(w):
    day = min(365, round((w - 0.5) * 7))
    return (dt.date(2025, 1, 1) + dt.timedelta(day - 1)).month
c_weeks = list(range(1, 53))
c_mean, c_lo, c_hi, c_cur, c_months = [], [], [], [], []
for wk in c_weeks:
    yr_vals = [clm[wk][y] for y in [2021, 2022, 2023, 2024, 2025] if y in clm[wk]]
    if yr_vals:
        c_mean.append(round(sum(yr_vals)/len(yr_vals))); c_lo.append(round(min(yr_vals))); c_hi.append(round(max(yr_vals)))
    else:
        c_mean.append(None); c_lo.append(None); c_hi.append(None)
    c_cur.append(round(cur_clm[wk]) if wk in cur_clm else None)
    c_months.append(week_month(wk))
_last_w = max(cur_clm) if cur_clm else None
DATA["claims_seasonal"] = {
    "weeks": c_weeks, "months": c_months,
    "mean": c_mean, "lo": c_lo, "hi": c_hi, "current": c_cur,
    "base_label": "2021–2025 avg",
    "last_week": _last_w,
    "last_date": cur_clm_date[_last_w] if _last_w else None,
}

# ------------------------------------------------- withheld employment taxes
# US Treasury DTS. Two eras, zero-gap splice:
#   Table IV (federal_tax_deposits), "Withheld Income and Employment Taxes"  2005-10-03 → 2023-02-13
#   Table II (deposits_withdrawals_operating_cash), "Taxes - Withheld
#   Individual/FICA"                                                          2023-02-14 → present
# Monthly sum of daily deposits; YoY vs same calendar month; 3mma of YoY.
# Validated against Variant Perception (Nov 19 2025 chart): Oct-2025 = +5.5%
# unsmoothed / +5.6% 3mma — exact match.
print("Treasury DTS withheld employment taxes...")
FD = "https://api.fiscaldata.treasury.gov/services/api/fiscal_service/v1/accounting/dts/"

def fiscaldata(endpoint, filt):
    rows, page = [], 1
    while True:
        r = requests.get(FD + endpoint,
                         params={"filter": filt, "page[size]": "10000", "page[number]": str(page)},
                         headers=HDR, timeout=120)
        r.raise_for_status()
        js = r.json()
        rows += js["data"]
        if page >= js["meta"]["total-pages"]:
            break
        page += 1
    return rows

wh_daily = defaultdict(float)
for x in fiscaldata("federal_tax_deposits",
                    "tax_deposit_type:eq:Withheld Income and Employment Taxes"):
    wh_daily[x["record_date"]] += float(x["tax_deposit_today_amt"])
for x in fiscaldata("deposits_withdrawals_operating_cash",
                    "transaction_catg:eq:Taxes - Withheld Individual/FICA,transaction_type:eq:Deposits"):
    wh_daily[x["record_date"]] += float(x["transaction_today_amt"])

wh_monthly = defaultdict(float)
for d, v in wh_daily.items():
    wh_monthly[d[:7]] += v
cur_month = dt.date.today().strftime("%Y-%m")
wh_months = sorted(m for m in wh_monthly if m < cur_month)   # complete months only

wh_yoy = OrderedDict()
for m in wh_months:
    prev = f"{int(m[:4])-1:04d}-{m[5:7]}"
    if prev in wh_monthly and wh_monthly[prev] > 0:
        wh_yoy[m] = (wh_monthly[m] / wh_monthly[prev] - 1) * 100
wh_keys = list(wh_yoy)
wh_raw, wh_3mma = [], []
for i, m in enumerate(wh_keys):
    if m < "2022-01":
        continue
    date = m + "-01"
    wh_raw.append([date, round(wh_yoy[m], 2)])
    wh_3mma.append([date, round(sum(wh_yoy[wh_keys[j]] for j in range(i - 2, i + 1)) / 3, 2)])
DATA["withheld_raw"] = wh_raw
DATA["withheld_3mma"] = wh_3mma

# ---------------------------------------------------------------- meta
DATA["meta"] = {"generated": dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%d %H:%M UTC")}

# ---------------------------------------------------------------- render
TEMPLATE = r'''<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>US High-Frequency Economic Data — Acheron Insights</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Space+Grotesk:wght@400;500;600;700&family=IBM+Plex+Sans:ital,wght@0,400;0,500;0,600;1,400&family=IBM+Plex+Mono:wght@400;500;600&display=swap" rel="stylesheet">
<script src="https://cdnjs.cloudflare.com/ajax/libs/Chart.js/4.4.1/chart.umd.min.js"></script>
<style>
:root{
  --paper:#F7F1E6; --card:#FCF8EF; --ink:#2A2520; --muted:#8A7E6E; --faint:#B6AB98;
  --teal:#0E756C; --teal-2:#12988B; --orange:#D2622A; --orange-2:#E0743B;
  --border:#E4DAC7; --border-strong:#D8CBB2; --grid:rgba(42,37,32,.06);
  --up:#0E756C; --down:#C0492A;
}
*{box-sizing:border-box}
html,body{margin:0;padding:0}
body{
  background:var(--paper);
  color:var(--ink);
  font-family:"IBM Plex Sans",system-ui,sans-serif;
  font-size:15px; line-height:1.5;
  -webkit-font-smoothing:antialiased;
  background-image:radial-gradient(circle at 1px 1px, rgba(42,37,32,.018) 1px, transparent 0);
  background-size:22px 22px;
}
.wrap{max-width:1360px; margin:0 auto; padding:34px 30px 60px}

/* ---------- header ---------- */
header.top{display:flex; justify-content:space-between; align-items:flex-end; gap:24px;
  border-bottom:2px solid var(--ink); padding-bottom:18px; margin-bottom:6px; flex-wrap:wrap}
.brand{display:flex; align-items:baseline; gap:12px}
.mark{font-family:"Space Grotesk"; font-weight:700; letter-spacing:.16em; font-size:12px;
  text-transform:uppercase; color:var(--orange)}
.mark .dot{color:var(--teal)}
h1{font-family:"Space Grotesk"; font-weight:600; font-size:30px; line-height:1.05;
  margin:14px 0 0; letter-spacing:-.01em}
.sub{color:var(--muted); font-size:13.5px; margin-top:7px; max-width:640px}
.stamp{text-align:right; font-family:"IBM Plex Mono"; font-size:11.5px; color:var(--muted);
  line-height:1.7; white-space:nowrap}
.stamp b{color:var(--ink); font-weight:600}

/* ---------- section ---------- */
.section-head{display:flex; align-items:center; gap:14px; margin:34px 0 16px}
.section-head .lbl{font-family:"Space Grotesk"; font-weight:600; font-size:12px;
  letter-spacing:.15em; text-transform:uppercase; color:var(--ink); white-space:nowrap}
.section-head .rule{height:1px; background:var(--border-strong); flex:1}
.section-head .cnt{font-family:"IBM Plex Mono"; font-size:11px; color:var(--faint)}

/* ---------- grid ---------- */
.grid{display:grid; grid-template-columns:repeat(auto-fill,minmax(258px,1fr)); gap:14px}
.card{background:var(--card); border:1px solid var(--border); border-radius:5px;
  padding:14px 15px 10px; display:flex; flex-direction:column; position:relative;
  box-shadow:0 1px 0 rgba(42,37,32,.03)}
.card .eyebrow{font-family:"Space Grotesk"; font-weight:600; font-size:9.5px; letter-spacing:.13em;
  text-transform:uppercase; margin-bottom:3px}
.card.real .eyebrow{color:var(--teal)}
.card.fin .eyebrow{color:var(--orange)}
.card h3{font-family:"Space Grotesk"; font-weight:500; font-size:14.5px; margin:0; letter-spacing:-.005em}
.card .desc{color:var(--muted); font-size:11px; margin-top:2px; min-height:14px}
.readout{display:flex; align-items:baseline; gap:8px; margin:9px 0 2px; flex-wrap:wrap}
.val{font-family:"IBM Plex Mono"; font-weight:600; font-size:24px; letter-spacing:-.02em; line-height:1}
.unit{font-size:11px; color:var(--muted); font-family:"IBM Plex Sans"}
.chip{font-family:"IBM Plex Mono"; font-size:11px; font-weight:500; padding:1px 6px; border-radius:3px;
  display:inline-flex; align-items:center; gap:3px}
.chip.up{color:var(--up); background:rgba(14,117,108,.10)}
.chip.down{color:var(--down); background:rgba(192,73,42,.10)}
.chip.flat{color:var(--muted); background:rgba(138,126,110,.12)}
.chip .win{color:var(--faint); font-weight:400; margin-left:1px}
.spark{position:relative; height:96px; margin:6px -4px 0}
.card .src{font-family:"IBM Plex Mono"; font-size:9px; color:var(--faint); margin-top:7px;
  padding-top:7px; border-top:1px solid var(--border); display:flex; justify-content:space-between}

/* ---------- seasonal ---------- */
.seasonal-grid{display:grid; grid-template-columns:repeat(auto-fit,minmax(460px,1fr)); gap:16px}
.card.wide{padding:16px 18px 12px}
.card.wide h3{font-size:17px; font-weight:600}
.legend{display:flex; gap:16px; align-items:center; margin:10px 0 2px; flex-wrap:wrap;
  font-size:11.5px; color:var(--muted)}
.legend .it{display:flex; align-items:center; gap:6px}
.legend .swatch{width:16px; height:3px; border-radius:2px}
.legend .swatch.band{height:11px; width:16px; border-radius:2px; opacity:.9}
.season-readout{display:flex; gap:26px; margin:8px 0 4px; flex-wrap:wrap}
.season-readout .blk .k{font-size:10.5px; color:var(--muted); font-family:"Space Grotesk";
  letter-spacing:.05em; text-transform:uppercase}
.season-readout .blk .v{font-family:"IBM Plex Mono"; font-weight:600; font-size:19px; margin-top:2px}
.spark.tall{height:230px}

footer{margin-top:40px; padding-top:18px; border-top:1px solid var(--border-strong);
  color:var(--muted); font-size:11.5px; line-height:1.65}
footer .notes{max-width:1000px}
footer b{color:var(--ink); font-weight:600}
footer .disc{margin-top:12px; font-size:10.5px; color:var(--faint); font-family:"IBM Plex Mono",monospace}

@media (max-width:640px){
  .wrap{padding:22px 16px 44px}
  h1{font-size:24px}
  .stamp{text-align:left}
  .seasonal-grid{grid-template-columns:1fr}
  .section-head{flex-wrap:wrap}
  .section-head .lbl{white-space:normal}
  .section-head .cnt{display:none}
}
@media (prefers-reduced-motion:reduce){*{animation:none!important; transition:none!important}}
</style>
</head>
<body>
<div class="wrap">

  <header class="top">
    <div>
      <div class="brand"><span class="mark">Acheron<span class="dot">·</span>Insights</span></div>
      <h1>US High-Frequency Economic Data</h1>
      <div class="sub">A real-time read on activity, labor, credit and prices from weekly and
        daily source data — refreshed straight from the primary agencies.</div>
    </div>
    <div class="stamp" id="stamp"></div>
  </header>

  <div class="section-head">
    <span class="lbl">High-Frequency Indicators</span>
    <span class="rule"></span>
    <span class="cnt" id="hf-count"></span>
  </div>
  <div class="grid" id="grid"></div>

  <div class="section-head">
    <span class="lbl">Seasonal Comparison — 2026 vs 5-Year Average</span>
    <span class="rule"></span>
    <span class="cnt">2021–2025 baseline</span>
  </div>
  <div class="seasonal-grid" id="seasonal"></div>

  <footer>
    <div class="notes">
      <b>Sources.</b> Federal Reserve Bank of Dallas (WEI), Atlanta Fed (GDPNow), Cleveland Fed
      (inflation nowcast), Chicago Fed (NFCI) and St.&nbsp;Louis Fed (financial stress) via FRED;
      US DOL / ETA jobless claims; Federal Reserve H.8 (bank loans); Freddie Mac PMMS (mortgage);
      Association of American Railroads via FRED (rail); US Census Bureau BFS (business applications);
      TSA (passenger throughput); US Treasury Daily Treasury Statement (withheld employment taxes).<br>
      <b>Notes.</b> Rail traffic is monthly (carloads + intermodal); the indicator panel uses the
      seasonally-adjusted series, the seasonal-comparison panel the non-adjusted series. The CPI
      nowcast is the Cleveland Fed's headline-CPI model chained to a year-over-year rate. Seasonal
      panels average the five most recent complete calendar years (2021–2025); shaded band shows the
      5-year min–max, lines are 7-day smoothed (TSA). GDPNow resets each quarter, producing the
      sawtooth pattern. Withheld employment taxes are the monthly sum of daily Treasury deposits (DTS Table IV spliced to the redesigned Table II at Feb 2023), shown as year-over-year change — teal line is the 3-month average, bars unsmoothed.
      <div class="disc">Redbook, Citi Economic Surprise and IBES forward EPS from the reference layout are
      proprietary subscription series without public APIs and are not reproduced here; the Cleveland Fed
      nowcast stands in for a public high-frequency inflation gauge. Data is provided for research and is
      not investment advice.</div>
    </div>
  </footer>

</div>

<script>
const DATA = {};

// ---- palette (explicit hex; CSS vars don't resolve inside canvas) ----
const C = {
  ink:"#2A2520", muted:"#8A7E6E", faint:"#B6AB98",
  teal:"#0E756C", tealFill:"rgba(14,117,108,0.13)",
  orange:"#D2622A", orangeFill:"rgba(210,98,42,0.13)",
  grid:"rgba(42,37,32,0.06)", band:"rgba(14,117,108,0.14)"
};
const RM = window.matchMedia("(prefers-reduced-motion: reduce)").matches;

// ---- formatting ----
function fmtNum(v, dp, suffix){
  if(v===null||v===undefined||isNaN(v)) return "–";
  let s = Math.abs(v).toLocaleString("en-US",{minimumFractionDigits:dp,maximumFractionDigits:dp});
  return (v<0?"−":"")+s+(suffix||"");
}
function fmtSigned(v, dp, suffix){
  if(v===null||v===undefined||isNaN(v)) return "–";
  const sign = v>0?"+":(v<0?"−":"");
  return sign+Math.abs(v).toLocaleString("en-US",{minimumFractionDigits:dp,maximumFractionDigits:dp})+(suffix||"");
}
function fmtDate(iso){
  const [y,m,d]=iso.split("-").map(Number);
  const mo=["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"][m-1];
  return d?`${mo} ${d}, ${y}`:`${mo} ${y}`;
}

// change over ~30 calendar days (uniform across frequencies)
function change30(series){
  if(series.length<2) return null;
  const last = series[series.length-1];
  const lastT = Date.parse(last[0]);
  let ref = series[0];
  for(let i=series.length-1;i>=0;i--){
    if(lastT - Date.parse(series[i][0]) >= 30*864e5){ ref = series[i]; break; }
  }
  return last[1]-ref[1];
}

// sparse year-boundary tick callback factory
function yearTicks(labels){
  return function(value, index){
    const lab = labels[index]; if(!lab) return "";
    const yr = lab.slice(0,4);
    if(index===0) return yr;
    return labels[index-1].slice(0,4)!==yr ? yr : "";
  };
}

function makeGradient(ctx, area, colorFill){
  const g = ctx.createLinearGradient(0, area.top, 0, area.bottom);
  g.addColorStop(0, colorFill);
  g.addColorStop(1, "rgba(0,0,0,0)");
  return g;
}

// ---- standard sparkline panel (optional faint bars behind the line) ----
function spark(canvas, series, accent, accentFill, dp, suffix, barSeries){
  const labels = series.map(p=>p[0]);
  const values = series.map(p=>p[1]);
  const ctx = canvas.getContext("2d");
  const ptR = values.map((_,i)=> i===values.length-1 ? 2.6 : 0);
  const datasets = [{
    label:"3mma",
    data:values, borderColor:accent, borderWidth:1.6, tension:.16, order:1,
    pointRadius:ptR, pointBackgroundColor:accent, pointBorderColor:accent,
    fill:barSeries?false:true,
    backgroundColor:(c)=>{const {ctx,chartArea}=c.chart; return chartArea?makeGradient(ctx,chartArea,accentFill):accentFill;}
  }];
  if(barSeries){
    const bmap = Object.fromEntries(barSeries);
    datasets.push({
      type:"bar", label:"unsmoothed", order:2,
      data:labels.map(l=> bmap[l]!==undefined? bmap[l] : null),
      backgroundColor:"rgba(42,37,32,0.32)", borderWidth:0,
      barPercentage:.6, categoryPercentage:.9
    });
  }
  return new Chart(ctx,{
    type:"line",
    data:{labels, datasets},
    options:{
      responsive:true, maintainAspectRatio:false, animation:RM?false:{duration:600},
      interaction:{intersect:false, mode:"index"},
      plugins:{legend:{display:false}, tooltip:{
        backgroundColor:"#2A2520", padding:8, cornerRadius:4, displayColors:false,
        titleFont:{family:"IBM Plex Mono", size:10}, bodyFont:{family:"IBM Plex Mono", size:11, weight:"600"},
        callbacks:{ title:(it)=>fmtDate(it[0].label),
          label:(it)=> (barSeries? it.dataset.label+": " : "") + fmtNum(it.raw,dp,suffix) }
      }},
      scales:{
        x:{grid:{display:false}, border:{color:C.grid},
           ticks:{autoSkip:false, maxRotation:0, font:{family:"IBM Plex Mono", size:9.5}, color:C.faint,
                  callback:yearTicks(labels), padding:2}},
        y:{position:"right", grid:{color:C.grid, drawTicks:false},
           border:{display:false},
           ticks:{maxTicksLimit:4, font:{family:"IBM Plex Mono", size:9.5}, color:C.muted, padding:4,
                  callback:(v)=>fmtNum(v,dp,"")}}
      }
    }
  });
}

// ---- panel configs ----
const PANELS = [
  {key:"wei",  fam:"real", cat:"Activity",  title:"Fed Weekly Economic Index",
   desc:"Composite scaled to y/y GDP growth", dp:2, suffix:"", src:"Dallas Fed · FRED"},
  {key:"gdpnow", fam:"real", cat:"Growth",   title:"Atlanta Fed GDPNow",
   desc:"Nowcast of current-quarter real GDP", dp:2, suffix:"%", src:"Atlanta Fed"},
  {key:"rail", fam:"real", cat:"Activity",   title:"US Rail Traffic",
   desc:"Carloads + intermodal, SA (monthly)", dp:0, suffix:"k", src:"AAR · FRED"},
  {key:"icsa", fam:"real", cat:"Labor",      title:"Initial Jobless Claims",
   desc:"New filings, seasonally adjusted", dp:0, suffix:"k", src:"US DOL · FRED"},
  {key:"ccsa", fam:"real", cat:"Labor",      title:"Continued Claims",
   desc:"Insured unemployment, SA", dp:2, suffix:"M", scale:0.001, src:"US DOL · FRED"},
  {key:"withheld_3mma", bars:"withheld_raw", fam:"real", cat:"Labor · Income",
   title:"Withheld Employment Taxes",
   desc:"Treasury withholding receipts, y/y", dp:1, suffix:"%", src:"US Treasury DTS"},
  {key:"loan_yoy", fam:"fin", cat:"Credit",  title:"Bank Loan Growth",
   desc:"Loans & leases, all banks", dp:2, suffix:"%", src:"Fed H.8 · FRED"},
  {key:"nfci", fam:"fin", cat:"Fin. Conditions", title:"Chicago Fed NFCI",
   desc:"Positive = tighter than average", dp:2, suffix:"", src:"Chicago Fed · FRED"},
  {key:"stress", fam:"fin", cat:"Fin. Conditions", title:"St. Louis Fed Stress",
   desc:"Positive = above-average stress", dp:2, suffix:"", src:"St. Louis Fed · FRED"},
  {key:"mortgage", fam:"fin", cat:"Credit",  title:"30-Yr Mortgage Rate",
   desc:"Freddie Mac 30-yr fixed survey", dp:2, suffix:"%", src:"Freddie Mac · FRED"},
  {key:"cpi_nowcast", fam:"fin", cat:"Inflation", title:"Cleveland Fed CPI Nowcast",
   desc:"Model nowcast, headline CPI y/y", dp:2, suffix:"%", src:"Cleveland Fed"}
];

function renderStandard(){
  const grid = document.getElementById("grid");
  document.getElementById("hf-count").textContent = PANELS.length + " series";
  PANELS.forEach(cfg=>{
    const series = DATA[cfg.key];
    if(!series || !series.length) return;
    const sc = cfg.scale || 1;
    const last = series[series.length-1];
    const lastVal = last[1]*sc;
    const chg = change30(series);
    const chgScaled = chg===null?null:chg*sc;
    const accent = cfg.fam==="real"?C.teal:C.orange;
    const accentFill = cfg.fam==="real"?C.tealFill:C.orangeFill;

    const card = document.createElement("div");
    card.className = "card "+cfg.fam;
    let chipCls="flat", arrow="→";
    if(chgScaled!==null && Math.abs(chg)>1e-9){ chipCls = chg>0?"up":"down"; arrow = chg>0?"▲":"▼"; }
    const chipTxt = chgScaled===null ? "" :
      `<span class="chip ${chipCls}">${arrow} ${fmtNum(Math.abs(chgScaled),cfg.dp,cfg.suffix)}<span class="win">30d</span></span>`;

    card.innerHTML = `
      <div class="eyebrow">${cfg.cat}</div>
      <h3>${cfg.title}</h3>
      <div class="desc">${cfg.desc}</div>
      <div class="readout">
        <span class="val">${fmtNum(lastVal,cfg.dp,"")}</span>
        <span class="unit">${cfg.suffix||""}</span>
        ${chipTxt}
      </div>
      <div class="spark"><canvas></canvas></div>
      <div class="src"><span>${cfg.src}</span><span>as of ${fmtDate(last[0])}</span></div>`;
    grid.appendChild(card);
    spark(card.querySelector("canvas"), series, accent, accentFill, cfg.dp, cfg.suffix,
          cfg.bars ? DATA[cfg.bars] : null);
  });
}

// ---- seasonal panel ----
function seasonalChart(canvas, axisLabels, band_lo, band_hi, mean, current, xTickCb){
  const ctx = canvas.getContext("2d");
  return new Chart(ctx,{
    type:"line",
    data:{labels:axisLabels, datasets:[
      // band: lo (invisible baseline) then hi filled to previous
      {data:band_lo, borderColor:"rgba(0,0,0,0)", pointRadius:0, fill:false, tension:.3, spanGaps:true},
      {data:band_hi, borderColor:"rgba(0,0,0,0)", pointRadius:0, fill:"-1",
       backgroundColor:C.band, tension:.3, spanGaps:true},
      {label:"5-yr avg", data:mean, borderColor:C.teal, borderWidth:1.8, pointRadius:0,
       tension:.3, borderDash:[5,3], fill:false, spanGaps:true},
      {label:"2026", data:current, borderColor:C.orange, borderWidth:2.4, pointRadius:0,
       tension:.25, fill:false, spanGaps:true}
    ]},
    options:{
      responsive:true, maintainAspectRatio:false, animation:RM?false:{duration:700},
      interaction:{intersect:false, mode:"index"},
      plugins:{legend:{display:false}, tooltip:{
        backgroundColor:"#2A2520", padding:9, cornerRadius:4,
        titleFont:{family:"IBM Plex Mono", size:10}, bodyFont:{family:"IBM Plex Mono", size:10.5},
        filter:(it)=> it.datasetIndex>=2,
        callbacks:{ title:(it)=>it[0].label,
          label:(it)=> `${it.dataset.label}: ${fmtNum(it.raw,0,"")}` }
      }},
      scales:{
        x:{grid:{display:false}, border:{color:C.grid},
           ticks:{autoSkip:false, maxRotation:0, font:{family:"IBM Plex Mono", size:10}, color:C.muted,
                  callback:xTickCb, padding:3}},
        y:{position:"right", grid:{color:C.grid, drawTicks:false}, border:{display:false},
           ticks:{maxTicksLimit:6, font:{family:"IBM Plex Mono", size:10}, color:C.muted, padding:4,
                  callback:(v)=> v>=1e6? (v/1e6).toFixed(1)+"M" : v>=1e3? (v/1e3).toFixed(0)+"k" : v}}
      }
    }
  });
}

const MON = ["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"];

function renderSeasonal(){
  const host = document.getElementById("seasonal");

  // -- TSA --
  const t = DATA.tsa_seasonal;
  if(t){
    // latest current vs seasonal mean at same day
    let li=-1; for(let i=0;i<t.current.length;i++) if(t.current[i]!==null) li=i;
    const curV = t.current[li], meanV = t.mean[li];
    const gap = meanV? (curV/meanV-1)*100 : null;
    const tickCb = function(value,index){
      const md = t.axis[index]; if(!md) return "";
      const [mm,dd]=md.split("-"); return dd==="01"? MON[parseInt(mm)-1] : "";
    };
    const card = document.createElement("div");
    card.className="card wide real";
    card.innerHTML = `
      <div class="eyebrow">Consumer · Air Travel</div>
      <h3>TSA Airport Throughput</h3>
      <div class="desc">Daily checkpoint passengers vs the 5-year seasonal norm</div>
      <div class="season-readout">
        <div class="blk"><div class="k">Latest (${fmtDate(t.last_date)})</div><div class="v" style="color:${C.orange}">${(curV/1e6).toFixed(2)}M</div></div>
        <div class="blk"><div class="k">5-yr avg, same day</div><div class="v" style="color:${C.teal}">${(meanV/1e6).toFixed(2)}M</div></div>
        <div class="blk"><div class="k">vs seasonal</div><div class="v" style="color:${gap>=0?C.up:C.down}">${fmtSigned(gap,1,"%")}</div></div>
      </div>
      <div class="legend">
        <div class="it"><span class="swatch" style="background:${C.orange}"></span>2026</div>
        <div class="it"><span class="swatch" style="background:${C.teal}"></span>${t.base_label}</div>
        <div class="it"><span class="swatch band" style="background:${C.band}"></span>5-yr range</div>
      </div>
      <div class="spark tall"><canvas></canvas></div>
      <div class="src"><span>TSA · daily throughput</span><span>7-day smoothed</span></div>`;
    host.appendChild(card);
    seasonalChart(card.querySelector("canvas"), t.axis, t.lo, t.hi, t.mean, t.current, tickCb);
  }

  // -- Business applications --
  const b = DATA.bizapps_seasonal;
  if(b){
    let li=-1; for(let i=0;i<b.current.length;i++) if(b.current[i]!==null) li=i;
    const curV=b.current[li], meanV=b.mean[li];
    const gap = meanV? (curV/meanV-1)*100 : null;
    const labels = b.weeks.map(w=>"W"+w);
    const tickCb = function(value,index){
      const m = b.months[index]; const prevM = index>0? b.months[index-1]:0;
      return (index===0 || m!==prevM) && m>0 ? MON[m-1] : "";
    };
    const card = document.createElement("div");
    card.className="card wide real";
    card.innerHTML = `
      <div class="eyebrow">Business Formation</div>
      <h3>Weekly Business Applications</h3>
      <div class="desc">New business applications (NSA) vs the 5-year seasonal norm</div>
      <div class="season-readout">
        <div class="blk"><div class="k">Latest (week ${b.last_week})</div><div class="v" style="color:${C.orange}">${(curV/1e3).toFixed(1)}k</div></div>
        <div class="blk"><div class="k">5-yr avg, same week</div><div class="v" style="color:${C.teal}">${(meanV/1e3).toFixed(1)}k</div></div>
        <div class="blk"><div class="k">vs seasonal</div><div class="v" style="color:${gap>=0?C.up:C.down}">${fmtSigned(gap,1,"%")}</div></div>
      </div>
      <div class="legend">
        <div class="it"><span class="swatch" style="background:${C.orange}"></span>2026</div>
        <div class="it"><span class="swatch" style="background:${C.teal}"></span>${b.base_label}</div>
        <div class="it"><span class="swatch band" style="background:${C.band}"></span>5-yr range</div>
      </div>
      <div class="spark tall"><canvas></canvas></div>
      <div class="src"><span>US Census · BFS weekly</span><span>not seasonally adjusted</span></div>`;
    host.appendChild(card);
    seasonalChart(card.querySelector("canvas"), labels, b.lo, b.hi, b.mean, b.current, tickCb);
  }

  // -- Rail traffic (monthly, NSA) --
  const r = DATA.rail_seasonal;
  if(r){
    let li=-1; for(let i=0;i<r.current.length;i++) if(r.current[i]!==null) li=i;
    const curV=r.current[li], meanV=r.mean[li];
    const gap = meanV? (curV/meanV-1)*100 : null;
    const labels = r.months.map(m=>MON[m-1]);
    const tickCb = function(value,index){ return MON[index]; };
    const card = document.createElement("div");
    card.className="card wide real";
    card.innerHTML = `
      <div class="eyebrow">Activity · Freight</div>
      <h3>US Rail Traffic</h3>
      <div class="desc">Carloads + intermodal (NSA) vs the 5-year seasonal norm</div>
      <div class="season-readout">
        <div class="blk"><div class="k">Latest (${MON[r.last_month-1]} 2026)</div><div class="v" style="color:${C.orange}">${(curV/1e6).toFixed(2)}M</div></div>
        <div class="blk"><div class="k">5-yr avg, same month</div><div class="v" style="color:${C.teal}">${(meanV/1e6).toFixed(2)}M</div></div>
        <div class="blk"><div class="k">vs seasonal</div><div class="v" style="color:${gap>=0?C.up:C.down}">${fmtSigned(gap,1,"%")}</div></div>
      </div>
      <div class="legend">
        <div class="it"><span class="swatch" style="background:${C.orange}"></span>2026</div>
        <div class="it"><span class="swatch" style="background:${C.teal}"></span>${r.base_label}</div>
        <div class="it"><span class="swatch band" style="background:${C.band}"></span>5-yr range</div>
      </div>
      <div class="spark tall"><canvas></canvas></div>
      <div class="src"><span>AAR · FRED (monthly)</span><span>not seasonally adjusted</span></div>`;
    host.appendChild(card);
    seasonalChart(card.querySelector("canvas"), labels, r.lo, r.hi, r.mean, r.current, tickCb);
  }

  // -- Initial jobless claims (weekly, NSA) --
  const cl = DATA.claims_seasonal;
  if(cl){
    let li=-1; for(let i=0;i<cl.current.length;i++) if(cl.current[i]!==null) li=i;
    const curV=cl.current[li], meanV=cl.mean[li];
    const gap = meanV? (curV/meanV-1)*100 : null;
    const labels = cl.weeks.map(w=>"W"+w);
    const tickCb = function(value,index){
      const m = cl.months[index]; const prevM = index>0? cl.months[index-1]:0;
      return (index===0 || m!==prevM) && m>0 ? MON[m-1] : "";
    };
    const card = document.createElement("div");
    card.className="card wide real";
    card.innerHTML = `
      <div class="eyebrow">Labor · Claims</div>
      <h3>Initial Jobless Claims</h3>
      <div class="desc">Weekly initial claims (NSA) vs the 5-year seasonal norm</div>
      <div class="season-readout">
        <div class="blk"><div class="k">Latest (${fmtDate(cl.last_date)})</div><div class="v" style="color:${C.orange}">${(curV/1e3).toFixed(0)}k</div></div>
        <div class="blk"><div class="k">5-yr avg, same week</div><div class="v" style="color:${C.teal}">${(meanV/1e3).toFixed(0)}k</div></div>
        <div class="blk"><div class="k">vs seasonal</div><div class="v" style="color:${gap<=0?C.up:C.down}">${fmtSigned(gap,1,"%")}</div></div>
      </div>
      <div class="legend">
        <div class="it"><span class="swatch" style="background:${C.orange}"></span>2026</div>
        <div class="it"><span class="swatch" style="background:${C.teal}"></span>${cl.base_label}</div>
        <div class="it"><span class="swatch band" style="background:${C.band}"></span>5-yr range</div>
      </div>
      <div class="spark tall"><canvas></canvas></div>
      <div class="src"><span>US DOL · FRED (weekly)</span><span>not seasonally adjusted</span></div>`;
    host.appendChild(card);
    seasonalChart(card.querySelector("canvas"), labels, cl.lo, cl.hi, cl.mean, cl.current, tickCb);
  }
}

// ---- stamp ----
function renderStamp(){
  const el = document.getElementById("stamp");
  el.innerHTML = `<div>Generated <b>${(DATA.meta&&DATA.meta.generated)||"—"}</b></div>`;
}

renderStamp();
renderStandard();
renderSeasonal();
</script>
</body>
</html>
'''

import re
payload = json.dumps(DATA)
html, n = re.subn(r"const DATA = \{\};\n",
                  lambda m: "const DATA = " + payload + ";\n",
                  TEMPLATE)
assert n == 1, f"expected exactly one DATA placeholder, found {n}"
OUT = "us_high_frequency_dashboard.html"
with open(OUT, "w") as f:
    f.write(html)
sizes = {k: (len(v) if isinstance(v, list) else "obj") for k, v in DATA.items()}
print("wrote", OUT, f"({len(html):,} bytes)")
for k, v in sizes.items(): print(f"  {k:20s}: {v}")
